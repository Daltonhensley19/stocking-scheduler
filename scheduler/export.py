"""Export the schedule as HTML (for print / PDF / PNG) and as an Excel workbook."""

from __future__ import annotations

from html import escape
from pathlib import Path

from .models import Team, Day, Associate, DayAssociate, working_associates, column_runs, minutes_to_label, ROLE_PICKER


def _date_title(day: Day) -> str:
    return day.date.strftime("%A, %B %d, %Y").replace(" 0", " ")


def _assoc_details(a: Associate, w: DayAssociate) -> list[str]:
    """Role / duty / aisles / shift lines (without the name)."""
    lines = [a.role + (" (optional)" if a.optional else "")]
    if w.duty:
        lines.append(w.duty)
    if w.aisles:
        lines.append(f"Topstock: {w.aisles}")
    if w.shift_start is not None and w.shift_end is not None:
        lines.append(f"{minutes_to_label(w.shift_start)}–{minutes_to_label(w.shift_end)}")
    return lines


def _assoc_header(a: Associate, w: DayAssociate) -> str:
    bits = [f"<b>{escape(a.name)}</b>"]
    for line in _assoc_details(a, w):
        bits.append(f"<span class='sub'>{escape(line)}</span>")
    return "<br>".join(bits)


# QTextDocument (used for print / PDF / PNG) supports a subset of CSS, so keep
# this simple: element/class selectors, colors, sizes, padding.
CSS = """
body { font-family: Helvetica, Arial, sans-serif; font-size: 9pt; color: #111; }
h1 { font-size: 15pt; margin: 0 0 2px 0; color: #0b2a4a; }
h2 { font-size: 11pt; margin: 0 0 8px 0; color: #333; font-weight: normal; }
th, td { border: 1px solid #9aa4b1; padding: 3px 4px; vertical-align: middle; }
th { background-color: #0b2a4a; color: white; font-weight: bold; }
th .sub { font-weight: normal; font-size: 8pt; color: #d9e2ec; }
td.time { background-color: #e9eef4; font-weight: bold; text-align: right; }
td.blk { text-align: center; }
td.blk .note { font-size: 8pt; color: #333; }
td.blk .rng { font-size: 7.5pt; color: #444; }
td.empty { background-color: #ffffff; }
p.footer { font-size: 8pt; color: #444; margin-top: 8px; }
p.sub { font-size: 10pt; color: #333; margin: 0 0 8px 0; }
"""

TABLE = "<table width='100%' cellspacing='0' cellpadding='3'>"


def board_html(team: Team, day: Day) -> str:
    """The whole-team grid for one day, one HTML table with merged blocks."""
    s = team.settings
    slots = s.slots()
    people = working_associates(team, day)

    col_w = f"{max(1, 94 // max(1, len(people)))}%"
    head = ["<th width='6%'>Time</th>"]
    for a, w in people:
        head.append(f"<th width='{col_w}'>{_assoc_header(a, w)}</th>")

    # Build per-column run lookup: row -> Run for the first row of each run,
    # and a set of rows covered by a rowspan (skipped when emitting cells).
    starts: list[dict[int, object]] = []
    covered: list[set[int]] = []
    for a, _ in people:
        runs = column_runs(day, a.id, slots, s.slot_minutes)
        starts.append({r.row: r for r in runs})
        covered.append({r.row + k for r in runs for k in range(1, r.rows)})

    rows_html = []
    for i, start in enumerate(slots):
        cells = [f"<td class='time'>{minutes_to_label(start)}</td>"]
        for col, (a, _) in enumerate(people):
            if i in covered[col]:
                continue
            run = starts[col].get(i)
            if run is None:
                cells.append("<td class='empty'></td>")
                continue
            seg = run.segment
            color = team.task_color(seg.task)
            inner = f"<b>{escape(seg.task)}</b>"
            if seg.note:
                inner += f"<br><span class='note'>{escape(seg.note)}</span>"
            inner += f"<br><span class='rng'>{minutes_to_label(seg.start)}–{minutes_to_label(seg.end)}</span>"
            cells.append(f"<td class='blk' rowspan='{run.rows}' bgcolor='{color}'>{inner}</td>")
        rows_html.append("<tr>" + "".join(cells) + "</tr>")

    footer = f"<p class='footer'>{escape(s.footer)}</p>" if s.footer else ""
    notes = f"<p class='footer'><b>Notes:</b> {escape(day.notes)}</p>" if day.notes.strip() else ""
    return (
        f"<html><head><style>{CSS}</style></head><body>"
        f"<h1>{escape(s.title)}</h1><h2>{_date_title(day)}</h2>"
        f"{TABLE}<tr>{''.join(head)}</tr>{''.join(rows_html)}</table>"
        f"{notes}{footer}</body></html>"
    )


def person_html(team: Team, day: Day, a: Associate, w: DayAssociate, standalone: bool = True) -> str:
    """One associate's day as a simple Time / Assignment list (like the PRINT tabs)."""
    s = team.settings
    rows = []
    for seg in sorted(day.segs(a.id), key=lambda x: x.start):
        color = team.task_color(seg.task)
        note = escape(seg.note) if seg.note else ""
        rows.append(
            f"<tr><td class='time'>{minutes_to_label(seg.start)} – {minutes_to_label(seg.end)}</td>"
            f"<td bgcolor='{color}'><b>{escape(seg.task)}</b></td><td>{note}</td></tr>"
        )
    if not rows:
        rows.append("<tr><td colspan='3'><i>Nothing assigned yet.</i></td></tr>")
    details = " &nbsp;·&nbsp; ".join(escape(x) for x in _assoc_details(a, w))
    body = (
        f"<h1>{escape(a.name)} — {_date_title(day)}</h1>"
        f"<p class='sub'>{details}</p>"
        f"{TABLE}<tr><th width='22%'>Time</th><th width='40%'>Assignment</th><th width='38%'>Note</th></tr>"
        f"{''.join(rows)}</table>"
    )
    if s.footer:
        body += f"<p class='footer'>{escape(s.footer)}</p>"
    if not standalone:
        return body
    return f"<html><head><style>{CSS}</style></head><body>{body}</body></html>"


def all_people_html(team: Team, day: Day) -> str:
    """Every working associate's personal schedule, one per page."""
    pages = [
        person_html(team, day, a, w, standalone=False) for a, w in working_associates(team, day)
    ]
    sep = "<div style='page-break-before: always'></div>"
    return f"<html><head><style>{CSS}</style></head><body>{sep.join(pages)}</body></html>"


# --- Excel --------------------------------------------------------------------


def export_xlsx(path: str | Path, team: Team, day: Day) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    s = team.settings
    slots = s.slots()
    people = working_associates(team, day)

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Schedule"

    thin = Side(style="thin", color="9AA4B1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hdr_fill = PatternFill("solid", fgColor="0B2A4A")
    hdr_font = Font(bold=True, color="FFFFFF")
    time_fill = PatternFill("solid", fgColor="E9EEF4")

    def fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_color.lstrip("#").upper())

    ncols = 1 + len(people)
    ws.cell(1, 1, s.title).font = Font(bold=True, size=14, color="0B2A4A")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 2))
    ws.cell(2, 1, _date_title(day)).font = Font(size=11)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(ncols, 2))

    # Header block: name / role+duty / aisles
    r0 = 4
    labels = ["Associate", "Role / Duty", "Topstock"]
    for k, lab in enumerate(labels):
        c = ws.cell(r0 + k, 1, lab)
        c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, center, border
    for col, (a, w) in enumerate(people, start=2):
        role = a.role + (" (opt)" if a.optional else "")
        vals = [a.name, f"{role}{' · ' + w.duty if w.duty else ''}", w.aisles]
        for k, v in enumerate(vals):
            c = ws.cell(r0 + k, col, v)
            c.fill, c.alignment, c.border = hdr_fill, center, border
            c.font = Font(bold=(k == 0), color="FFFFFF", size=11 if k == 0 else 9)
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions["A"].width = 10

    first = r0 + len(labels)
    for i, start in enumerate(slots):
        c = ws.cell(first + i, 1, minutes_to_label(start))
        c.fill, c.alignment, c.border = time_fill, Alignment(horizontal="right", vertical="center"), border
        c.font = Font(bold=True)
        for col in range(2, ncols + 1):
            ws.cell(first + i, col).border = border

    for col, (a, _) in enumerate(people, start=2):
        for run in column_runs(day, a.id, slots, s.slot_minutes):
            seg = run.segment
            top = first + run.row
            text = seg.task
            if seg.note:
                text += f"\n{seg.note}"
            text += f"\n{minutes_to_label(seg.start)}–{minutes_to_label(seg.end)}"
            c = ws.cell(top, col, text)
            c.fill, c.alignment, c.border = fill(team.task_color(seg.task)), center, border
            c.font = Font(size=9)
            if run.rows > 1:
                ws.merge_cells(start_row=top, start_column=col, end_row=top + run.rows - 1, end_column=col)
                for k in range(run.rows):
                    ws.cell(top + k, col).border = border

    note_row = first + len(slots) + 1
    if day.notes.strip():
        ws.cell(note_row, 1, "Notes: " + day.notes).font = Font(italic=True, size=9)
        note_row += 1
    if s.footer:
        ws.cell(note_row, 1, s.footer).font = Font(italic=True, size=9, color="444444")
    ws.freeze_panes = ws.cell(first, 2)

    # One sheet per associate (mirrors the PRINT tabs in the reference sheet)
    used: set[str] = set()
    for a, w in people:
        title = a.name[:28] or "Associate"
        base, n = title, 2
        while title in used or title == "Daily Schedule":
            title = f"{base} {n}"
            n += 1
        used.add(title)
        p = wb.create_sheet(title)
        p.cell(1, 1, f"{a.name} — {_date_title(day)}").font = Font(bold=True, size=13, color="0B2A4A")
        p.merge_cells("A1:C1")
        sub = a.role + (" (optional)" if a.optional else "")
        if w.duty:
            sub += f" · {w.duty}"
        if w.aisles:
            sub += f" · Topstock {w.aisles}"
        p.cell(2, 1, sub)
        p.merge_cells("A2:C2")
        for k, lab in enumerate(["Time", "Assignment", "Note"], start=1):
            c = p.cell(4, k, lab)
            c.fill, c.font, c.alignment, c.border = hdr_fill, hdr_font, center, border
        r = 5
        for seg in sorted(day.segs(a.id), key=lambda x: x.start):
            t = p.cell(r, 1, f"{minutes_to_label(seg.start)} – {minutes_to_label(seg.end)}")
            t.border, t.fill = border, time_fill
            c = p.cell(r, 2, seg.task)
            c.fill, c.border, c.font = fill(team.task_color(seg.task)), border, Font(bold=True)
            nn = p.cell(r, 3, seg.note)
            nn.border = border
            r += 1
        if a.role == ROLE_PICKER:
            p.cell(r + 1, 1, "Tip: put the case count or helper's name in the Note column.").font = Font(
                italic=True, size=9, color="666666"
            )
        p.column_dimensions["A"].width = 18
        p.column_dimensions["B"].width = 32
        p.column_dimensions["C"].width = 30

    wb.save(str(path))
